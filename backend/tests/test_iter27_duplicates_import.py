"""Iter27: Duplicate detection in import (preview/validate/import), required column enforcement,
template xlsx styling, GET /question-bank/duplicates, POST /questions/bulk-delete."""
import io
import os
import csv
import pytest
import requests
import openpyxl

def _read_env():
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    return line.split("=", 1)[1].strip()
    except Exception:
        return ""
    return ""

BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or _read_env()).rstrip("/")
API = f"{BASE_URL}/api"
ADMIN = {"email": "admin@lms.id", "password": "Admin@12345"}


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=ADMIN, timeout=20)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def tryout(admin):
    payload = {"title": "TEST_ITER27_DUP", "description": "iter27", "subject": "Matematika",
               "duration_minutes": 30, "price": 0, "published": False, "kind": "standalone"}
    r = admin.post(f"{API}/admin/tryouts", json=payload, timeout=20)
    assert r.status_code in (200, 201), r.text
    tid = r.json()["id"]
    yield tid
    admin.delete(f"{API}/admin/tryouts/{tid}", timeout=20)


HEADERS = ["type", "text", "option_a", "option_b", "option_c", "option_d", "option_e",
           "correct", "points", "competency", "tingkat"]


def _csv(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=HEADERS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")


def _row(text="Soal uji 2+2=?", correct="B"):
    return {"type": "single", "text": text, "option_a": "3", "option_b": "4",
            "option_c": "5", "option_d": "6", "option_e": "",
            "correct": correct, "points": 10, "competency": "numerasi", "tingkat": "mudah"}


# ---- Preview duplicate marking: within-file + existing-in-tryout ----
def test_preview_within_file_duplicate(admin, tryout):
    rows = [_row("Soal A"), _row("Soal B"), _row("Soal A")]  # row 3 duplicate of row 1
    files = {"file": ("in.csv", _csv(rows), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["valid"] == 2
    assert data["duplicates"] == 1
    assert data["invalid"] == 0
    dup_row = data["rows"][2]
    assert dup_row["ok"] is False
    assert dup_row["duplicate"] and "Sama dengan baris 2" in dup_row["duplicate"]
    assert dup_row["error"] is None  # not counted as invalid


def test_preview_existing_in_tryout_duplicate(admin, tryout):
    # Insert one question first via import-rows
    base = _row("Soal yang sudah ada", correct="A")
    r = admin.post(f"{API}/admin/tryouts/{tryout}/questions/import-rows", json={"rows": [base]}, timeout=20)
    assert r.status_code == 200 and r.json()["imported"] == 1

    # Preview a file containing same question + a fresh one
    rows = [_row("Soal yang sudah ada", correct="A"), _row("Soal baru X")]
    files = {"file": ("in.csv", _csv(rows), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["duplicates"] == 1
    assert data["valid"] == 1
    assert data["rows"][0]["duplicate"] and "sudah ada" in data["rows"][0]["duplicate"]


def test_validate_endpoint_marks_duplicates(admin, tryout):
    rows = [_row("Soal yang sudah ada", correct="A"), _row("Soal baru Y")]
    r = admin.post(f"{API}/admin/questions/import/validate?tryout_id={tryout}",
                   json={"rows": rows}, timeout=20)
    assert r.status_code == 200
    d = r.json()
    assert d["duplicates"] == 1 and d["valid"] == 1


# ---- Import endpoints skip duplicates ----
def test_import_rows_skips_duplicates(admin, tryout):
    rows = [_row("Soal yang sudah ada", correct="A"), _row("Soal baru Z")]
    r = admin.post(f"{API}/admin/tryouts/{tryout}/questions/import-rows",
                   json={"rows": rows}, timeout=20)
    assert r.status_code == 200
    d = r.json()
    assert d["imported"] == 1
    assert len(d["duplicates"]) == 1
    assert "sudah ada" in d["duplicates"][0]


def test_legacy_import_skips_duplicates(admin, tryout):
    rows = [_row("Soal yang sudah ada", correct="A"), _row("Soal baru legacy")]
    files = {"file": ("in.csv", _csv(rows), "text/csv")}
    r = admin.post(f"{API}/admin/tryouts/{tryout}/questions/import", files=files, timeout=30)
    assert r.status_code == 200
    d = r.json()
    assert d["imported"] == 1
    assert len(d["duplicates"]) == 1


# ---- Missing required column => 400 ----
def test_missing_required_column_400(admin, tryout):
    bad_headers = [h for h in HEADERS if h != "option_e"]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=bad_headers, extrasaction="ignore")
    w.writeheader()
    w.writerow({k: _row()[k] for k in bad_headers})
    files = {"file": ("bad.csv", ("\ufeff" + buf.getvalue()).encode("utf-8"), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=20)
    assert r.status_code == 400
    assert "Kolom wajib tidak ada" in r.json().get("detail", "")
    assert "option_e" in r.json()["detail"]


def test_missing_points_row_error(admin, tryout):
    row = _row("Row empty points")
    row["points"] = ""
    files = {"file": ("in.csv", _csv([row]), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=20)
    assert r.status_code == 200
    d = r.json()
    assert d["invalid"] == 1
    assert "points" in (d["rows"][0]["error"] or "").lower()


def test_missing_competency_row_error(admin, tryout):
    row = _row("Row empty comp")
    row["competency"] = ""
    files = {"file": ("in.csv", _csv([row]), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=20)
    d = r.json()
    assert d["invalid"] == 1
    assert "competency" in (d["rows"][0]["error"] or "").lower()


def test_missing_tingkat_row_error(admin, tryout):
    row = _row("Row empty tingkat")
    row["tingkat"] = ""
    files = {"file": ("in.csv", _csv([row]), "text/csv")}
    r = admin.post(f"{API}/admin/questions/import/preview?tryout_id={tryout}", files=files, timeout=20)
    d = r.json()
    assert d["invalid"] == 1
    assert "tingkat" in (d["rows"][0]["error"] or "").lower()


# ---- Template xlsx: required headers styled blue + Petunjuk WAJIB rows ----
def test_template_xlsx_styling_and_guide(admin):
    r = admin.get(f"{API}/admin/questions/template.xlsx", timeout=30)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Soal" in wb.sheetnames and "Petunjuk" in wb.sheetnames
    ws = wb["Soal"]
    required = {"type", "text", "option_a", "option_b", "option_c", "option_d", "option_e",
                "correct", "points", "competency", "tingkat"}
    for cell in ws[1]:
        if cell.value in required:
            # 0E7490 blue teal fill
            fg = (cell.fill.fgColor.rgb or "").upper()
            assert "0E7490" in fg, f"required col {cell.value} missing blue fill (got {fg})"
    info = wb["Petunjuk"]
    tags = [row[0].value for row in info.iter_rows() if row and row[0].value]
    assert "WAJIB" in tags
    assert "OPSIONAL" in tags
    assert "Duplikat" in tags


# ---- GET /question-bank/duplicates + bulk-delete ----
def test_duplicates_endpoint_and_bulk_delete(admin):
    # Create a fresh tryout and insert 2 identical questions via POST /questions (manual, allowed to duplicate)
    r = admin.post(f"{API}/admin/tryouts", json={
        "title": "TEST_ITER27_DUP2", "description": "d", "subject": "Matematika",
        "duration_minutes": 20, "price": 0, "published": False, "kind": "standalone"}, timeout=20)
    tid = r.json()["id"]
    try:
        text = "DupText_Iter27_Unique_XYZ"
        qpayload = {"type": "single", "text": text,
                    "options": [{"id": "o1", "text": "A"}, {"id": "o2", "text": "B"}],
                    "correct_answers": ["o1"], "points": 5,
                    "competency": "numerasi", "difficulty_label": "mudah"}
        ids = []
        for _ in range(2):
            rr = admin.post(f"{API}/admin/tryouts/{tid}/questions", json=qpayload, timeout=20)
            assert rr.status_code in (200, 201), rr.text
            ids.append(rr.json()["id"])

        r = admin.get(f"{API}/admin/question-bank/duplicates", timeout=30)
        assert r.status_code == 200
        data = r.json()
        found = next((g for g in data["groups"] if g["text"] == text), None)
        assert found is not None
        assert found["count"] == 2
        assert found["same_tryout"] is True
        # has_explanation / tryout_has_results / tryout_title present
        for it in found["items"]:
            assert "has_explanation" in it and "tryout_has_results" in it and "tryout_title" in it

        # bulk-delete one of them
        rr = admin.post(f"{API}/admin/questions/bulk-delete",
                        json={"question_ids": [ids[0]]}, timeout=20)
        assert rr.status_code == 200
        assert rr.json()["deleted"] == 1

        # Verify removed from duplicates
        r = admin.get(f"{API}/admin/question-bank/duplicates", timeout=30)
        found2 = next((g for g in r.json()["groups"] if g["text"] == text), None)
        assert found2 is None or found2["count"] < 2
    finally:
        admin.delete(f"{API}/admin/tryouts/{tid}", timeout=20)


def test_bulk_delete_empty_400(admin):
    r = admin.post(f"{API}/admin/questions/bulk-delete", json={"question_ids": []}, timeout=20)
    assert r.status_code == 400
