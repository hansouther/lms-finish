from fastapi import APIRouter, HTTPException, Depends, BackgroundTasks, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel, EmailStr
from typing import List, Optional
import io
import csv
import re
import zipfile
import asyncio
import openpyxl

from database import db
from utils import new_id, now_iso
from security import require_roles, hash_password
from emailer import notify_new_tryout, notify_bid_accepted, notify_new_material
from storage import put_object, MIME_TYPES, APP_NAME
from routes_files import save_question_image, IMAGE_EXT
import notifications
import analysis
import irt
from grading import grade_attempt
from tryout_session import schedule_status, remaining_seconds, answered_count

router = APIRouter(prefix="/api/admin", tags=["admin"])
admin_only = require_roles("admin")


async def _active_student_recipients():
    enrolls = await db.enrollments.find({}, {"_id": 0, "student_id": 1}).to_list(5000)
    ids = list({e["student_id"] for e in enrolls})
    students = await db.users.find({"id": {"$in": ids}, "role": "student"}, {"_id": 0, "email": 1, "name": 1}).to_list(5000)
    return [{"email": s["email"], "name": s.get("name")} for s in students if s.get("email")]


# ---------- Models ----------
class NewsBody(BaseModel):
    title: str
    content: str
    category: str = "Pengumuman"
    published: bool = True


class CalendarBody(BaseModel):
    title: str
    description: Optional[str] = None
    date: str
    type: str = "academic"


class CourseBody(BaseModel):
    title: str
    description: str
    subject: str
    level: str = "Umum"
    price: int = 0
    thumbnail: Optional[str] = None
    active: bool = True


class SessionBody(BaseModel):
    id: Optional[str] = None
    date: str
    start_time: str
    end_time: str
    topic: Optional[str] = None


class SlotBody(BaseModel):
    title: str
    subject: str
    required_qualifications: List[str] = []
    course_id: Optional[str] = None
    open_to_all: bool = False
    notes: Optional[str] = None
    sessions: List[SessionBody] = []
    # legacy single-session support
    date: Optional[str] = None
    start_time: Optional[str] = None
    end_time: Optional[str] = None


class TryoutBody(BaseModel):
    title: str
    description: Optional[str] = None
    subject: str
    duration_minutes: int = 60
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    published: bool = False
    course_id: Optional[str] = None
    kind: str = "standalone"
    irt_scale: str = irt.DEFAULT_SCALE  # snbt | tka | raw | custom
    irt_scale_custom: Optional[dict] = None  # {min, max, mean, sd} bila custom
    price: int = 0  # harga Try Out (Rupiah); 0 = gratis


class IrtScaleBody(BaseModel):
    irt_scale: str
    irt_scale_custom: Optional[dict] = None


def _clean_scale(key, custom):
    key = (key or irt.DEFAULT_SCALE).lower()
    if key not in irt.SCALE_PRESETS:
        raise HTTPException(status_code=400, detail="Skala skor IRT tidak valid")
    if key == "custom":
        custom, err = irt.validate_custom_scale(custom or {})
        if err:
            raise HTTPException(status_code=400, detail=err)
        return key, custom
    return key, None


async def _rescore_attempts(tryout_id: str, tryout: dict):
    """Hitung ulang theta/skor IRT semua attempt tersubmit memakai skala Try Out."""
    scale = irt.resolve_scale(tryout)
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    attempts = await db.attempts.find({"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0}).to_list(10000)
    for a in attempts:
        res = irt.score_attempt(questions, a.get("per_question"), scale)
        await db.attempts.update_one({"id": a["id"]}, {"$set": res})
    return len(attempts)


class QuestionBody(BaseModel):
    type: str  # single | multiple | truefalse | truefalse_multi | essay
    text: str
    image_url: Optional[str] = None
    tf_labels: Optional[List[str]] = None
    options: List[dict] = []
    correct_answers: List[str] = []
    points: int = 1
    order: int = 0
    competency: str = "umum"  # numerasi | literasi | umum (AKM)
    difficulty_label: Optional[str] = None  # mudah | sedang | sulit (label manual untuk IRT)


class BulkCompetencyBody(BaseModel):
    question_ids: List[str]
    competency: str  # numerasi | literasi | umum


class BulkDifficultyBody(BaseModel):
    question_ids: List[str]
    difficulty_label: str  # mudah | sedang | sulit


class BroadcastBody(BaseModel):
    title: str
    message: str
    priority: str = "normal"


class UserBody(BaseModel):
    name: str
    email: EmailStr
    password: str
    role: str  # tutor | proctor | student
    phone: Optional[str] = None
    school_id: Optional[str] = None
    qualifications: List[str] = []


class UserUpdateBody(BaseModel):
    role: Optional[str] = None
    school_id: Optional[str] = None
    status: Optional[str] = None  # pending | approved | rejected


class SchoolBody(BaseModel):
    name: str
    city: Optional[str] = None


# ---------- Dashboard ----------
@router.get("/stats")
async def stats(user: dict = Depends(admin_only)):
    return {
        "students": await db.users.count_documents({"role": "student"}),
        "tutors": await db.users.count_documents({"role": "tutor"}),
        "proctors": await db.users.count_documents({"role": "proctor"}),
        "courses": await db.courses.count_documents({}),
        "tryouts": await db.tryouts.count_documents({}),
        "open_slots": await db.teaching_slots.count_documents({"status": "open"}),
        "pending_partnerships": await db.partnerships.count_documents({"status": "new"}),
        "pending_verifications": await db.users.count_documents({"status": "pending"}),
        "total_attempts": await db.attempts.count_documents({"status": "submitted"}),
    }


@router.get("/pending-count")
async def pending_count(user: dict = Depends(admin_only)):
    return {"count": await db.users.count_documents({"status": "pending"})}


# ---------- News ----------
@router.get("/news")
async def list_news(user: dict = Depends(admin_only)):
    return await db.news.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.post("/news")
async def create_news(body: NewsBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "author": user["name"], "created_at": now_iso()}
    await db.news.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/news/{item_id}")
async def update_news(item_id: str, body: NewsBody, user: dict = Depends(admin_only)):
    res = await db.news.update_one({"id": item_id}, {"$set": body.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Berita tidak ditemukan")
    return {"ok": True}


@router.delete("/news/{item_id}")
async def delete_news(item_id: str, user: dict = Depends(admin_only)):
    await db.news.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Calendar ----------
@router.get("/calendar")
async def list_calendar(user: dict = Depends(admin_only)):
    return await db.calendar_events.find({}, {"_id": 0}).sort("date", 1).to_list(500)


@router.post("/calendar")
async def create_calendar(body: CalendarBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.calendar_events.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/calendar/{item_id}")
async def delete_calendar(item_id: str, user: dict = Depends(admin_only)):
    await db.calendar_events.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Courses ----------
@router.get("/courses")
async def list_courses(user: dict = Depends(admin_only)):
    items = await db.courses.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for c in items:
        c["enrolled_count"] = await db.enrollments.count_documents({"course_id": c["id"]})
    return items


@router.post("/courses")
async def create_course(body: CourseBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.courses.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/courses/{item_id}")
async def update_course(item_id: str, body: CourseBody, user: dict = Depends(admin_only)):
    res = await db.courses.update_one({"id": item_id}, {"$set": body.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    return {"ok": True}


@router.delete("/courses/{item_id}")
async def delete_course(item_id: str, user: dict = Depends(admin_only)):
    await db.courses.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Partnerships ----------
@router.get("/partnerships")
async def list_partnerships(user: dict = Depends(admin_only)):
    return await db.partnerships.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.put("/partnerships/{item_id}")
async def update_partnership(item_id: str, status: str, user: dict = Depends(admin_only)):
    await db.partnerships.update_one({"id": item_id}, {"$set": {"status": status}})
    return {"ok": True}


# ---------- Teaching Slots & Bidding ----------
@router.get("/slots")
async def list_slots(user: dict = Depends(admin_only)):
    slots = await db.teaching_slots.find({}, {"_id": 0}).sort("date", 1).to_list(300)
    for s in slots:
        s["bid_count"] = await db.bids.count_documents({"slot_id": s["id"]})
    return slots


def _build_sessions(body):
    sessions = []
    for i, s in enumerate(body.sessions):
        sessions.append({
            "id": s.id or new_id(),
            "no": i + 1,
            "date": s.date,
            "start_time": s.start_time,
            "end_time": s.end_time,
            "topic": s.topic,
        })
    if not sessions and body.date:
        sessions = [{
            "id": new_id(), "no": 1, "date": body.date,
            "start_time": body.start_time, "end_time": body.end_time, "topic": body.notes,
        }]
    return sessions


@router.post("/slots")
async def create_slot(body: SlotBody, user: dict = Depends(admin_only)):
    sessions = _build_sessions(body)
    if not sessions:
        raise HTTPException(status_code=400, detail="Minimal satu pertemuan diperlukan")
    first = sessions[0]
    doc = {
        "id": new_id(),
        "title": body.title,
        "subject": body.subject,
        "required_qualifications": body.required_qualifications,
        "course_id": body.course_id,
        "open_to_all": body.open_to_all,
        "notes": body.notes,
        "sessions": sessions,
        "date": first["date"],
        "start_time": first["start_time"],
        "end_time": first["end_time"],
        "status": "open",
        "tutor_id": None,
        "created_by": user["id"],
        "created_at": now_iso(),
    }
    await db.teaching_slots.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/slots/{slot_id}")
async def update_slot(slot_id: str, body: SlotBody, user: dict = Depends(admin_only)):
    slot = await db.teaching_slots.find_one({"id": slot_id})
    if not slot:
        raise HTTPException(status_code=404, detail="Slot tidak ditemukan")
    sessions = _build_sessions(body)
    if not sessions:
        raise HTTPException(status_code=400, detail="Minimal satu pertemuan diperlukan")
    first = sessions[0]
    old_count = len(slot.get("sessions") or [])
    updates = {
        "title": body.title, "subject": body.subject,
        "required_qualifications": body.required_qualifications,
        "course_id": body.course_id, "open_to_all": body.open_to_all, "notes": body.notes,
        "sessions": sessions, "date": first["date"],
        "start_time": first["start_time"], "end_time": first["end_time"],
    }
    await db.teaching_slots.update_one({"id": slot_id}, {"$set": updates})
    # Notify students if new sessions were added to a confirmed class
    if slot.get("status") == "confirmed" and len(sessions) > old_count:
        added = len(sessions) - old_count
        sids = await notifications.course_student_ids(body.course_id)
        if sids:
            await notifications.push(
                sids, "session",
                f"{added} pertemuan baru ditambahkan",
                f'Kelas "{body.title}" kini memiliki {len(sessions)} pertemuan',
                "/student/schedule",
            )
            students = await db.users.find({"id": {"$in": sids}}, {"_id": 0, "email": 1, "name": 1}).to_list(2000)
            await notify_new_material(students, body.title, f"{added} pertemuan baru ditambahkan ke jadwal kelas")
    return {"ok": True}


@router.delete("/slots/{slot_id}")
async def delete_slot(slot_id: str, user: dict = Depends(admin_only)):
    await db.teaching_slots.delete_one({"id": slot_id})
    await db.bids.delete_many({"slot_id": slot_id})
    await db.class_materials.delete_many({"slot_id": slot_id})
    await db.attendance.delete_many({"slot_id": slot_id})
    return {"ok": True}


@router.get("/slots/{slot_id}/bids")
async def slot_bids(slot_id: str, user: dict = Depends(admin_only)):
    return await db.bids.find({"slot_id": slot_id}, {"_id": 0}).sort("created_at", 1).to_list(200)


@router.post("/slots/{slot_id}/assign/{bid_id}")
async def assign_slot(slot_id: str, bid_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    bid = await db.bids.find_one({"id": bid_id, "slot_id": slot_id}, {"_id": 0})
    if not bid:
        raise HTTPException(status_code=404, detail="Bid tidak ditemukan")
    await db.teaching_slots.update_one(
        {"id": slot_id}, {"$set": {"status": "confirmed", "tutor_id": bid["tutor_id"]}}
    )
    await db.bids.update_one({"id": bid_id}, {"$set": {"status": "accepted"}})
    await db.bids.update_many(
        {"slot_id": slot_id, "id": {"$ne": bid_id}}, {"$set": {"status": "rejected"}}
    )
    slot = await db.teaching_slots.find_one({"id": slot_id}, {"_id": 0})
    tutor = await db.users.find_one({"id": bid["tutor_id"]}, {"_id": 0, "email": 1, "name": 1})
    if slot and tutor and tutor.get("email"):
        background.add_task(
            notify_bid_accepted, tutor["email"], tutor.get("name", "Tentor"),
            slot.get("title", ""), slot.get("date", ""),
            f'{slot.get("start_time", "")}-{slot.get("end_time", "")}',
        )
    return {"ok": True}


# ---------- Try Out & Questions ----------
@router.get("/tryouts")
async def list_tryouts(user: dict = Depends(admin_only)):
    items = await db.tryouts.find({"kind": {"$ne": "exercise"}}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for t in items:
        t["question_count"] = await db.questions.count_documents({"tryout_id": t["id"]})
        t["attempt_count"] = await db.attempts.count_documents({"tryout_id": t["id"], "status": "submitted"})
        t["in_progress_count"] = await db.attempts.count_documents({"tryout_id": t["id"], "status": "in_progress"})
        t["schedule_status"] = schedule_status(t)
    return items


@router.post("/tryouts")
async def create_tryout(body: TryoutBody, background: BackgroundTasks, user: dict = Depends(admin_only)):
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    doc = {"id": new_id(), **body.model_dump(), "irt_scale": key, "irt_scale_custom": custom,
           "created_by": user["id"], "created_at": now_iso()}
    await db.tryouts.insert_one(doc)
    if doc.get("published") and doc.get("kind", "standalone") == "standalone":
        recipients = await _active_student_recipients()
        if recipients:
            background.add_task(notify_new_tryout, recipients, doc["title"], doc["subject"])
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/tryouts/{tryout_id}")
async def update_tryout(tryout_id: str, body: TryoutBody, background: BackgroundTasks, user: dict = Depends(admin_only)):
    existing = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    updates = {**body.model_dump(), "irt_scale": key, "irt_scale_custom": custom}
    await db.tryouts.update_one({"id": tryout_id}, {"$set": updates})
    scale_changed = (existing.get("irt_scale") or irt.DEFAULT_SCALE) != key or (existing.get("irt_scale_custom") or None) != custom
    rescored = await _rescore_attempts(tryout_id, {**existing, **updates}) if scale_changed else 0
    if body.published and not existing.get("published") and body.kind == "standalone":
        recipients = await _active_student_recipients()
        if recipients:
            background.add_task(notify_new_tryout, recipients, body.title, body.subject)
    return {"ok": True, "rescored": rescored}


@router.get("/irt-scales")
async def irt_scales(user: dict = Depends(admin_only)):
    """Daftar preset skala skor IRT (SNBT, TKA, SD/SMP, Kustom)."""
    return list(irt.SCALE_PRESETS.values())


@router.get("/tryouts/{tryout_id}/irt-scale")
async def get_irt_scale(tryout_id: str, user: dict = Depends(admin_only)):
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    return {"irt_scale": t.get("irt_scale") or irt.DEFAULT_SCALE,
            "irt_scale_custom": t.get("irt_scale_custom"),
            "effective": irt.resolve_scale(t)}


@router.put("/tryouts/{tryout_id}/irt-scale")
async def set_irt_scale(tryout_id: str, body: IrtScaleBody, user: dict = Depends(admin_only)):
    """Ubah skala skor IRT sebuah Try Out lalu hitung ulang skor semua attempt tersubmit."""
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    key, custom = _clean_scale(body.irt_scale, body.irt_scale_custom)
    await db.tryouts.update_one({"id": tryout_id}, {"$set": {"irt_scale": key, "irt_scale_custom": custom}})
    t.update({"irt_scale": key, "irt_scale_custom": custom})
    rescored = await _rescore_attempts(tryout_id, t)
    return {"ok": True, "irt_scale": key, "irt_scale_custom": custom,
            "effective": irt.resolve_scale(t), "rescored": rescored}


@router.delete("/tryouts/{tryout_id}")
async def delete_tryout(tryout_id: str, user: dict = Depends(admin_only)):
    await db.tryouts.delete_one({"id": tryout_id})
    await db.questions.delete_many({"tryout_id": tryout_id})
    return {"ok": True}


@router.get("/tryouts/{tryout_id}/questions")
async def list_questions(tryout_id: str, user: dict = Depends(admin_only)):
    return await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).sort("order", 1).to_list(300)


@router.post("/tryouts/{tryout_id}/questions")
async def create_question(tryout_id: str, body: QuestionBody, user: dict = Depends(admin_only)):
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    count = await db.questions.count_documents({"tryout_id": tryout_id})
    payload = body.model_dump()
    dl = (payload.get("difficulty_label") or "").lower()
    payload["difficulty_label"] = dl if dl in irt.DIFFICULTY_LABELS else None
    doc = {"id": new_id(), "tryout_id": tryout_id, **payload,
           "irt_a": None, "irt_b": None, "irt_c": None,
           "irt_calibrated": False, "p_value": None, "point_biserial": None,
           "response_count": 0}
    if not doc.get("order"):
        doc["order"] = count + 1
    await db.questions.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.put("/questions/{question_id}")
async def update_question(question_id: str, body: QuestionBody, user: dict = Depends(admin_only)):
    payload = body.model_dump()
    dl = (payload.get("difficulty_label") or "").lower()
    payload["difficulty_label"] = dl if dl in irt.DIFFICULTY_LABELS else None
    res = await db.questions.update_one({"id": question_id}, {"$set": payload})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Soal tidak ditemukan")
    return {"ok": True}


@router.delete("/questions/{question_id}")
async def delete_question(question_id: str, user: dict = Depends(admin_only)):
    await db.questions.delete_one({"id": question_id})
    return {"ok": True}


@router.post("/tryouts/{tryout_id}/questions/bulk-competency")
async def bulk_competency(tryout_id: str, body: BulkCompetencyBody, user: dict = Depends(admin_only)):
    comp = body.competency if body.competency in ("numerasi", "literasi", "umum") else "umum"
    if not body.question_ids:
        raise HTTPException(status_code=400, detail="Tidak ada soal dipilih")
    res = await db.questions.update_many(
        {"tryout_id": tryout_id, "id": {"$in": body.question_ids}},
        {"$set": {"competency": comp}},
    )
    return {"updated": res.modified_count}


@router.post("/tryouts/{tryout_id}/questions/bulk-difficulty")
async def bulk_difficulty(tryout_id: str, body: BulkDifficultyBody, user: dict = Depends(admin_only)):
    lvl = (body.difficulty_label or "").lower()
    if lvl not in irt.DIFFICULTY_LABELS:
        raise HTTPException(status_code=400, detail="Tingkat kesulitan tidak valid")
    if not body.question_ids:
        raise HTTPException(status_code=400, detail="Tidak ada soal dipilih")
    res = await db.questions.update_many(
        {"tryout_id": tryout_id, "id": {"$in": body.question_ids}},
        {"$set": {"difficulty_label": lvl}},
    )
    return {"updated": res.modified_count}


@router.post("/tryouts/{tryout_id}/calibrate")
async def calibrate_tryout(tryout_id: str, user: dict = Depends(admin_only)):
    """Kalibrasi parameter IRT (3PL) tiap butir dari data jawaban peserta,
    lalu hitung ulang skor IR/theta untuk semua attempt tersubmit."""
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    tryout = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    scale = irt.resolve_scale(tryout)
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    attempts = await db.attempts.find(
        {"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0}
    ).to_list(10000)

    params = irt.calibrate(questions, attempts)
    calibrated = 0
    for qid, p in params.items():
        await db.questions.update_one({"id": qid}, {"$set": p})
        if p.get("irt_calibrated"):
            calibrated += 1

    # Terapkan parameter baru ke objek soal in-memory lalu skor ulang theta tiap attempt
    qmap = {q["id"]: q for q in questions}
    for qid, p in params.items():
        if qid in qmap:
            qmap[qid].update(p)
    rescored = 0
    for a in attempts:
        res = irt.score_attempt(list(qmap.values()), a.get("per_question"), scale)
        await db.attempts.update_one({"id": a["id"]}, {"$set": res})
        rescored += 1

    return {
        "calibrated": calibrated,
        "items_with_responses": len(params),
        "total_questions": len(questions),
        "attempts_scored": rescored,
        "min_responses_required": 5,
    }


@router.get("/tryouts/{tryout_id}/irt-blueprint")
async def irt_blueprint(tryout_id: str, user: dict = Depends(admin_only)):
    """Rekomendasi komposisi soal (Mudah/Sedang/Sulit) + Test Information Function
    untuk memenuhi kebutuhan penilaian IRT saat membuat Try Out."""
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    return irt.build_blueprint(questions)


@router.get("/tryouts/{tryout_id}/results")
async def tryout_results(tryout_id: str, user: dict = Depends(admin_only)):
    attempts = await db.attempts.find(
        {"tryout_id": tryout_id, "status": "submitted"}, {"_id": 0, "per_question": 0}
    ).to_list(2000)
    sids = list({a["student_id"] for a in attempts})
    students = await db.users.find({"id": {"$in": sids}}, {"_id": 0, "id": 1, "name": 1, "email": 1}).to_list(2000)
    smap = {s["id"]: s for s in students}
    for a in attempts:
        a["student"] = smap.get(a["student_id"])
    return sorted(attempts, key=lambda a: a.get("percentage", 0), reverse=True)


@router.get("/tryouts/{tryout_id}/monitor")
async def tryout_monitor(tryout_id: str, user: dict = Depends(admin_only)):
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    total_q = await db.questions.count_documents({"tryout_id": tryout_id})
    attempts = await db.attempts.find({"tryout_id": tryout_id}, {"_id": 0, "per_question": 0}).to_list(5000)
    grants = await db.tryout_retakes.find({"tryout_id": tryout_id, "used": False}, {"_id": 0}).to_list(5000)
    sids = list({a["student_id"] for a in attempts} | {g["student_id"] for g in grants})
    smap = {u["id"]: u for u in await db.users.find({"id": {"$in": sids}}, {"_id": 0, "id": 1, "name": 1, "email": 1}).to_list(5000)}
    rows = []
    for a in attempts:
        answers = a.get("answers") or {}
        status = a["status"]
        left = remaining_seconds(a, t) if status == "in_progress" else 0
        if status == "in_progress" and left <= 0:
            status = "expired"
        rows.append({
            "attempt_id": a["id"], "student": smap.get(a["student_id"]), "student_id": a["student_id"],
            "status": status, "answered_count": a.get("answered_count", answered_count(answers)),
            "total_questions": total_q, "started_at": a.get("started_at"), "submitted_at": a.get("submitted_at"),
            "last_saved_at": a.get("last_saved_at") or (a.get("submitted_at") if a["status"] == "submitted" else None),
            "remaining_seconds": left, "percentage": a.get("percentage") if a["status"] == "submitted" else None,
            "retake": bool(a.get("bypass_schedule")),
        })
    for g in grants:
        rows.append({"attempt_id": None, "student": smap.get(g["student_id"]), "student_id": g["student_id"],
                     "status": "retake_pending", "answered_count": 0, "total_questions": total_q,
                     "started_at": None, "submitted_at": None, "last_saved_at": None, "remaining_seconds": 0,
                     "percentage": None, "retake": True, "granted_at": g.get("created_at")})
    order = {"in_progress": 0, "expired": 1, "retake_pending": 2, "submitted": 3}
    rows.sort(key=lambda r: (order.get(r["status"], 9), (r["student"] or {}).get("name") or ""))
    summary = {k: sum(1 for r in rows if r["status"] == k) for k in order}
    return {"tryout": {"id": t["id"], "title": t["title"], "start_at": t.get("start_at"), "end_at": t.get("end_at"),
                       "duration_minutes": t.get("duration_minutes"), "schedule_status": schedule_status(t)},
            "summary": summary, "rows": rows, "server_now": now_iso()}


@router.post("/tryouts/{tryout_id}/attempts/{attempt_id}/reset")
async def reset_attempt(tryout_id: str, attempt_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    a = await db.attempts.find_one({"id": attempt_id, "tryout_id": tryout_id}, {"_id": 0, "student_id": 1, "status": 1})
    if not a:
        raise HTTPException(status_code=404, detail="Data pengerjaan tidak ditemukan")
    sid = a["student_id"]
    await db.attempts.delete_many({"tryout_id": tryout_id, "student_id": sid})
    await db.tryout_retakes.delete_many({"tryout_id": tryout_id, "student_id": sid, "used": False})
    await db.tryout_retakes.insert_one({"id": new_id(), "tryout_id": tryout_id, "student_id": sid, "used": False,
                                        "granted_by": user["id"], "created_at": now_iso()})
    if a["status"] == "submitted":
        from routes_student import _recalibrate_tryout
        background.add_task(_recalibrate_tryout, tryout_id)
    return {"ok": True}


@router.delete("/tryouts/{tryout_id}/retakes/{student_id}")
async def cancel_retake(tryout_id: str, student_id: str, user: dict = Depends(admin_only)):
    await db.tryout_retakes.delete_many({"tryout_id": tryout_id, "student_id": student_id, "used": False})
    return {"ok": True}


@router.post("/tryouts/{tryout_id}/attempts/{attempt_id}/force-submit")
async def force_submit(tryout_id: str, attempt_id: str, background: BackgroundTasks, user: dict = Depends(admin_only)):
    a = await db.attempts.find_one({"id": attempt_id, "tryout_id": tryout_id}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Data pengerjaan tidak ditemukan")
    if a["status"] == "submitted":
        raise HTTPException(status_code=400, detail="Sudah dikumpulkan")
    t = await db.tryouts.find_one({"id": tryout_id}, {"_id": 0})
    questions = await db.questions.find({"tryout_id": tryout_id}, {"_id": 0}).to_list(1000)
    result = grade_attempt(questions, a.get("answers") or {})
    irt_score = irt.score_attempt(questions, result.get("per_question"), irt.resolve_scale(t))
    await db.attempts.update_one({"id": attempt_id}, {"$set": {
        "status": "submitted", "submitted_at": now_iso(), "force_submitted_by": user["id"], **result, **irt_score}})
    from routes_student import _recalibrate_tryout
    background.add_task(_recalibrate_tryout, tryout_id)
    return {"ok": True, "percentage": result["percentage"]}


# ---------- Weakness Analysis & Score Export (all schools) ----------
@router.get("/analysis")
async def analysis_report(user: dict = Depends(admin_only)):
    return await analysis.build_report(None)


@router.get("/analysis/scores.csv")
async def analysis_scores_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.scores_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=nilai_siswa.csv"})


@router.get("/analysis/recap.csv")
async def analysis_recap_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.recap_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=rekap_kelemahan_siswa.csv"})


@router.get("/analysis/items.csv")
async def analysis_items_csv(user: dict = Depends(admin_only)):
    rep = await analysis.build_report(None)
    return Response(content=analysis.items_csv(rep), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=analisis_butir_soal.csv"})


# ---------- Broadcasts ----------
@router.get("/broadcasts")
async def list_broadcasts(user: dict = Depends(admin_only)):
    return await db.broadcasts.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.post("/broadcasts")
async def create_broadcast(body: BroadcastBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "author": user["name"], "created_at": now_iso()}
    await db.broadcasts.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/broadcasts/{item_id}")
async def delete_broadcast(item_id: str, user: dict = Depends(admin_only)):
    await db.broadcasts.delete_one({"id": item_id})
    return {"ok": True}


# ---------- Users & Schools ----------
@router.get("/users")
async def list_users(role: Optional[str] = None, user: dict = Depends(admin_only)):
    q = {"role": role} if role else {}
    items = await db.users.find(q, {"_id": 0, "password_hash": 0}).sort("created_at", -1).to_list(500)
    schools = await db.schools.find({}, {"_id": 0}).to_list(200)
    smap = {s["id"]: s["name"] for s in schools}
    for u in items:
        u["school_name"] = smap.get(u.get("school_id"))
    return items


@router.post("/users")
async def create_user(body: UserBody, user: dict = Depends(admin_only)):
    if body.role not in ("tutor", "proctor", "student"):
        raise HTTPException(status_code=400, detail="Peran tidak valid")
    email = body.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {
        "id": new_id(),
        "email": email,
        "password_hash": hash_password(body.password),
        "name": body.name,
        "role": body.role,
        "status": "approved",
        "phone": body.phone,
        "school_id": body.school_id,
        "qualifications": body.qualifications,
        "picture": None,
        "auth_provider": "password",
        "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    return {"ok": True, "id": doc["id"]}


@router.put("/users/{user_id}")
async def update_user(user_id: str, body: UserUpdateBody, user: dict = Depends(admin_only)):
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    updates = {}
    if body.role is not None:
        if body.role not in ("admin", "tutor", "proctor", "student"):
            raise HTTPException(status_code=400, detail="Peran tidak valid")
        updates["role"] = body.role
    if body.status is not None:
        if body.status not in ("pending", "approved", "rejected"):
            raise HTTPException(status_code=400, detail="Status tidak valid")
        updates["status"] = body.status
    if body.school_id is not None:
        updates["school_id"] = body.school_id or None
    if not updates:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    await db.users.update_one({"id": user_id}, {"$set": updates})
    return {"ok": True, **updates}


@router.delete("/users/{user_id}")
async def delete_user(user_id: str, user: dict = Depends(admin_only)):
    if user_id == user["id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus akun sendiri")
    await db.users.delete_one({"id": user_id})
    return {"ok": True}


@router.get("/schools")
async def list_schools(user: dict = Depends(admin_only)):
    items = await db.schools.find({}, {"_id": 0}).sort("name", 1).to_list(200)
    for s in items:
        s["student_count"] = await db.users.count_documents({"role": "student", "school_id": s["id"]})
    return items


@router.post("/schools")
async def create_school(body: SchoolBody, user: dict = Depends(admin_only)):
    doc = {"id": new_id(), **body.model_dump(), "created_at": now_iso()}
    await db.schools.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


# ---------- Course Content: Lessons & Exercises ----------
class LessonBody(BaseModel):
    title: str
    description: Optional[str] = None
    video_type: str = "youtube"  # youtube | upload
    video_url: Optional[str] = None
    attachments: List[dict] = []  # [{name, url, content_type}]
    order: int = 0


class ExerciseBody(BaseModel):
    title: str
    subject: Optional[str] = "Latihan"
    duration_minutes: int = 20
    description: Optional[str] = None


@router.get("/courses/{course_id}/content")
async def course_content(course_id: str, user: dict = Depends(admin_only)):
    course = await db.courses.find_one({"id": course_id}, {"_id": 0})
    if not course:
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    lessons = await db.lessons.find({"course_id": course_id}, {"_id": 0}).sort("order", 1).to_list(200)
    exercises = await db.tryouts.find({"course_id": course_id, "kind": "exercise"}, {"_id": 0}).sort("created_at", 1).to_list(200)
    for ex in exercises:
        ex["question_count"] = await db.questions.count_documents({"tryout_id": ex["id"]})
        ex["attempt_count"] = await db.attempts.count_documents({"tryout_id": ex["id"], "status": "submitted"})
    return {"course": course, "lessons": lessons, "exercises": exercises}


@router.post("/courses/{course_id}/lessons")
async def create_lesson(course_id: str, body: LessonBody, user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    count = await db.lessons.count_documents({"course_id": course_id})
    doc = {"id": new_id(), "course_id": course_id, **body.model_dump(), "created_at": now_iso()}
    if not doc.get("order"):
        doc["order"] = count + 1
    await db.lessons.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


@router.delete("/lessons/{lesson_id}")
async def delete_lesson(lesson_id: str, user: dict = Depends(admin_only)):
    await db.lessons.delete_one({"id": lesson_id})
    return {"ok": True}


@router.post("/courses/{course_id}/exercises")
async def create_exercise(course_id: str, body: ExerciseBody, user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    doc = {
        "id": new_id(), "title": body.title, "description": body.description,
        "subject": body.subject or "Latihan", "duration_minutes": body.duration_minutes,
        "start_at": now_iso(), "end_at": None, "published": True,
        "course_id": course_id, "kind": "exercise",
        "created_by": user["id"], "created_at": now_iso(),
    }
    await db.tryouts.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


# ---------- Bulk Import Questions (CSV / Excel) ----------
_LETTERS = ["a", "b", "c", "d", "e", "f"]


def _row_to_question(row: dict, order: int):
    t = str(row.get("type") or "").strip().lower()
    t = {"bs_majemuk": "truefalse_multi", "benar_salah_majemuk": "truefalse_multi", "truefalse-multi": "truefalse_multi"}.get(t, t)
    if t not in ("single", "multiple", "truefalse", "truefalse_multi", "essay"):
        return None, f"Tipe tidak valid: {t!r}"
    text = str(row.get("text") or "").strip()
    if not text:
        return None, "Kolom 'text' kosong"
    try:
        points = int(float(row.get("points") or 1))
    except Exception:
        points = 1
    options, correct = [], []
    for i, l in enumerate(_LETTERS):
        v = row.get(f"option_{l}")
        v = str(v).strip() if v is not None else ""
        if v:
            options.append({"id": f"o{i + 1}", "text": v})
    correct_raw = str(row.get("correct") or "").strip()
    if t in ("single", "multiple"):
        for tok in re.split(r"[;,|]", correct_raw):
            tok = tok.strip().lower()
            if not tok:
                continue
            if tok in _LETTERS:
                correct.append(f"o{_LETTERS.index(tok) + 1}")
            else:
                for o in options:
                    if o["text"].strip().lower() == tok:
                        correct.append(o["id"])
        if not options or not correct:
            return None, "Opsi atau kunci tidak lengkap"
    elif t == "truefalse":
        val = correct_raw.strip().lower()
        if val in ("true", "benar", "b", "ya", "1"):
            val = "true"
        elif val in ("false", "salah", "s", "tidak", "0"):
            val = "false"
        else:
            return None, "Kunci benar/salah tidak valid"
        correct, options = [val], []
    elif t == "truefalse_multi":
        keys = [k.strip().lower() for k in re.split(r"[;,|]", correct_raw) if k.strip()]
        if not options or len(keys) != len(options):
            return None, "Jumlah kunci (mis. B;S;B) harus sama dengan jumlah pernyataan di option_a.."
        for o, k in zip(options, keys):
            if k in ("true", "benar", "b", "ya", "1"):
                correct.append(f"{o['id']}:true")
            elif k in ("false", "salah", "s", "tidak", "0"):
                correct.append(f"{o['id']}:false")
            else:
                return None, f"Kunci benar/salah tidak valid: {k!r}"
    else:  # essay
        correct = [c.strip() for c in re.split(r"[|;]", correct_raw) if c.strip()]
        options = []
        if not correct:
            return None, "Kunci esai kosong"
    comp = str(row.get("competency") or row.get("kompetensi") or "umum").strip().lower()
    if comp not in ("numerasi", "literasi"):
        comp = "umum"
    dl = str(row.get("difficulty_label") or row.get("tingkat") or row.get("kesulitan") or "").strip().lower()
    _MAP = {"mudah": "mudah", "easy": "mudah", "sedang": "sedang", "medium": "sedang",
            "sulit": "sulit", "susah": "sulit", "hard": "sulit"}
    dl = _MAP.get(dl)
    lt = str(row.get("label_benar") or "").strip()
    lf = str(row.get("label_salah") or "").strip()
    tf_labels = [lt or "Benar", lf or "Salah"] if t in ("truefalse", "truefalse_multi") and (lt or lf) else None
    return {"type": t, "text": text, "options": options, "correct_answers": correct, "tf_labels": tf_labels,
            "points": points, "order": order, "competency": comp, "difficulty_label": dl}, None


def _parse_sheet(data: bytes, fname: str):
    rows = []
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        headers = None
        for r in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(c).strip().lower() if c is not None else "" for c in r]
                continue
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        for r in csv.DictReader(io.StringIO(text)):
            rows.append({(k or "").strip().lower(): v for k, v in r.items()})
    return rows


def _parse_zip(data: bytes):
    images, sheet, sheet_name = {}, None, ""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            name = info.filename
            base = name.replace("\\", "/").split("/")[-1]
            if info.is_dir() or not base or name.startswith("__MACOSX") or base.startswith("."):
                continue
            low = base.lower()
            if low.endswith((".xlsx", ".xls", ".csv")) and sheet is None:
                sheet, sheet_name = zf.read(info), low
            elif low.rsplit(".", 1)[-1] in IMAGE_EXT:
                images[low] = (base, zf.read(info))
    if sheet is None:
        raise ValueError("ZIP harus berisi 1 file Excel (.xlsx) atau CSV")
    return _parse_sheet(sheet, sheet_name), images


@router.post("/tryouts/{tryout_id}/questions/import")
async def import_questions(tryout_id: str, file: UploadFile = File(...), user: dict = Depends(admin_only)):
    if not await db.tryouts.find_one({"id": tryout_id}):
        raise HTTPException(status_code=404, detail="Try Out tidak ditemukan")
    data = await file.read()
    fname = (file.filename or "").lower()
    images = {}
    try:
        if fname.endswith(".zip"):
            rows, images = _parse_zip(data)
        else:
            rows = _parse_sheet(data, fname)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca berkas: {e}")

    start = await db.questions.count_documents({"tryout_id": tryout_id})
    imported, errors = 0, []
    for idx, row in enumerate(rows):
        q, err = _row_to_question(row, start + imported + 1)
        if err:
            errors.append(f"Baris {idx + 2}: {err}")
            continue
        img = str(row.get("gambar") or row.get("image") or row.get("image_url") or "").strip()
        q["image_url"] = None
        if img:
            if img.startswith(("https://", "http://", "/api/")):
                q["image_url"] = img
            elif img.split("/")[-1].lower() in images:
                base, blob = images[img.split("/")[-1].lower()]
                try:
                    q["image_url"] = await save_question_image(base, blob, user["id"])
                except Exception as e:
                    errors.append(f"Baris {idx + 2}: gambar {img!r} gagal diunggah ({e}), soal tetap diimpor tanpa gambar")
            else:
                errors.append(f"Baris {idx + 2}: gambar {img!r} tidak ditemukan (unggah sebagai ZIP berisi Excel + gambar), soal tetap diimpor tanpa gambar")
        q["id"] = new_id()
        q["tryout_id"] = tryout_id
        q.update({"irt_a": None, "irt_b": None, "irt_c": None, "irt_calibrated": False,
                  "p_value": None, "point_biserial": None, "response_count": 0})
        await db.questions.insert_one(q)
        imported += 1
    return {"imported": imported, "errors": errors}


@router.get("/questions/template")
async def questions_template(user: dict = Depends(admin_only)):
    csv_text = (
        "type,text,option_a,option_b,option_c,option_d,correct,points,competency,tingkat,gambar,label_benar,label_salah\n"
        "single,\"Berapa hasil 2+2?\",3,4,5,6,B,10,numerasi,mudah,,,\n"
        "single,\"Hasil dari $\\frac{1}{2} + \\frac{1}{3}$ adalah ...\",$\\frac{5}{6}$,$\\frac{2}{5}$,$\\frac{1}{6}$,$\\frac{3}{5}$,A,10,numerasi,sedang,,,\n"
        "multiple,\"Pilih bilangan genap\",2,3,4,5,\"A;C\",10,numerasi,sedang,,,\n"
        "single,\"Perhatikan gambar. Luas segitiga tersebut adalah ...\",$12\\text{ cm}^2$,$24\\text{ cm}^2$,$6\\text{ cm}^2$,$18\\text{ cm}^2$,A,10,numerasi,sedang,segitiga.png,,\n"
        "truefalse,\"Bumi berbentuk bulat\",,,,,benar,10,literasi,mudah,,,\n"
        "truefalse_multi,\"Bacalah teks berikut. Rani membeli 3 buku seharga Rp5.000 per buku dan membayar Rp20.000. Tentukan benar atau salah pernyataan berikut.\",Total harga buku Rp15.000,Uang kembalian Rp10.000,Jika membeli 4 buku uangnya pas,,\"B;S;B\",15,numerasi,sedang,,Sesuai,Tidak Sesuai\n"
        "essay,\"Ibu kota Indonesia?\",,,,,\"jakarta|dki jakarta\",10,literasi,sulit,,,\n"
    )
    return Response(content=csv_text, media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=template_soal.csv"})


@router.get("/questions/template.xlsx")
async def questions_template_xlsx(user: dict = Depends(admin_only)):
    """Template Excel untuk pembuatan soal skala besar (banyak soal sekaligus)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Soal"
    headers = ["type", "text", "option_a", "option_b", "option_c", "option_d",
               "option_e", "correct", "points", "competency", "tingkat", "gambar", "label_benar", "label_salah"]
    ws.append(headers)
    samples = [
        ["single", "Berapa hasil 2+2?", "3", "4", "5", "6", "", "B", 10, "numerasi", "mudah", ""],
        ["single", r"Hasil dari $\frac{1}{2} + \frac{1}{3}$ adalah ...", r"$\frac{5}{6}$", r"$\frac{2}{5}$", r"$\frac{1}{6}$", r"$\frac{3}{5}$", "", "A", 10, "numerasi", "sedang", ""],
        ["single", r"Akar-akar dari $$x^2 - 5x + 6 = 0$$ adalah ...", r"$x=2$ dan $x=3$", r"$x=-2$ dan $x=-3$", r"$x=1$ dan $x=6$", r"$x=-1$ dan $x=6$", "", "A", 10, "numerasi", "sulit", ""],
        ["single", "Perhatikan gambar. Luas segitiga tersebut adalah ...", r"$12\text{ cm}^2$", r"$24\text{ cm}^2$", r"$6\text{ cm}^2$", r"$18\text{ cm}^2$", "", "A", 10, "numerasi", "sedang", "segitiga.png"],
        ["multiple", "Pilih bilangan genap", "2", "3", "4", "5", "", "A;C", 10, "numerasi", "sedang", ""],
        ["truefalse", "Bumi berbentuk bulat", "", "", "", "", "", "benar", 10, "literasi", "mudah", ""],
        ["truefalse_multi", "Bacalah teks berikut. Rani membeli 3 buku seharga Rp5.000 per buku dan membayar Rp20.000. Tentukan benar atau salah setiap pernyataan.", "Total harga buku Rp15.000", "Uang kembalian Rp10.000", "Jika membeli 4 buku uangnya pas", "", "", "B;S;B", 15, "numerasi", "sedang", "", "Sesuai", "Tidak Sesuai"],
        ["truefalse_multi", r"Diketahui $f(x) = 2x + 3$. Tentukan benar/salah pernyataan berikut.", r"$f(1) = 5$", r"$f(0) = 2$", r"$f(-1) = 1$", "", "", "B;S;B", 15, "numerasi", "sulit", "", "", ""],
        ["essay", "Ibu kota Indonesia?", "", "", "", "", "", "jakarta|dki jakarta", 10, "literasi", "sulit", ""],
    ]
    for row in samples:
        ws.append(row)
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 48)

    info = wb.create_sheet("Petunjuk")
    guide = [
        ["Petunjuk Pengisian Template Soal (Impor Massal)"],
        [""],
        ["Kolom", "Keterangan"],
        ["type", "Tipe soal: single (PG), multiple (PG kompleks), truefalse (1 pernyataan), truefalse_multi (Benar/Salah majemuk: 1 teks + beberapa pernyataan), essay"],
        ["text", "Teks pertanyaan"],
        ["option_a..e", "Pilihan jawaban (untuk single/multiple). Kosongkan untuk truefalse/essay"],
        ["correct", "Kunci: huruf A-E (pisah ; untuk multiple), benar/salah untuk truefalse, B;S;B (urut sesuai option_a..e) untuk truefalse_multi, atau teks (pisah |) untuk essay"],
        ["truefalse_multi", "Tulis teks deskripsi/stimulus di kolom text, lalu tiap pernyataan di option_a, option_b, dst. Nilai dihitung proporsional dari jumlah pernyataan yang benar."],
        ["label_benar / label_salah", "OPSIONAL untuk truefalse & truefalse_multi. Ganti teks tombol, mis. Sesuai / Tidak Sesuai, Fakta / Opini, Ya / Tidak. Kosong = Benar / Salah."],
        ["points", "Poin/bobot soal (angka)"],
        ["competency", "numerasi | literasi | umum (untuk laporan AKM)"],
        ["tingkat", "Tingkat kesulitan IRT: mudah | sedang | sulit"],
        ["gambar", "OPSIONAL. Isi nama file gambar (mis. segitiga.png) lalu unggah ZIP berisi file Excel ini + semua gambarnya. Bisa juga diisi link https://... Kosongkan bila soal tanpa gambar."],
        [""],
        ["Rumus Matematika (LaTeX)", "Tulis rumus di antara tanda $...$ (sebaris) atau $$...$$ (baris sendiri). Berlaku di kolom text & option."],
        ["Contoh pecahan", r"$\frac{a}{b}$"],
        ["Contoh pangkat & akar", r"$x^2 + \sqrt{y} = \sqrt[3]{27}$"],
        ["Contoh integral & limit", r"$$\int_0^1 x^2\,dx = \frac{1}{3}$$   $\lim_{x \to 0} \frac{\sin x}{x} = 1$"],
        ["Contoh simbol", r"$\alpha, \beta, \pi, \theta, \le, \ge, \ne, \times, \div, \pm, \infty$"],
        ["Contoh matriks", r"$\begin{pmatrix} 1 & 2 \\ 3 & 4 \end{pmatrix}$"],
        ["Tanda dolar biasa", r"Tulis \$ untuk menampilkan simbol $ (mis. harga \$5)"],
        [""],
        ["Rekomendasi komposisi IRT", "±30% Mudah, 40% Sedang, 30% Sulit"],
        ["Catatan", "Parameter IRT (a,b,c) akan dikalibrasi otomatis dari hasil jawaban peserta."],
    ]
    for row in guide:
        info.append(row)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_soal.xlsx"},
    )


# ---------- Bulk Import Course Materials (ZIP) ----------
_VIDEO_EXT = {"mp4", "webm", "mov", "m4v"}
_DOC_EXT = {"pdf", "doc", "docx", "ppt", "pptx"}


@router.post("/courses/{course_id}/lessons/import-zip")
async def import_zip(course_id: str, file: UploadFile = File(...), user: dict = Depends(admin_only)):
    if not await db.courses.find_one({"id": course_id}):
        raise HTTPException(status_code=404, detail="Kursus tidak ditemukan")
    data = await file.read()
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except Exception:
        raise HTTPException(status_code=400, detail="Berkas ZIP tidak valid")

    created, skipped, errors = 0, 0, []
    order = await db.lessons.count_documents({"course_id": course_id})
    for name in zf.namelist():
        base = name.split("/")[-1]
        if not base or base.startswith(".") or "__MACOSX" in name:
            continue
        ext = base.rsplit(".", 1)[-1].lower() if "." in base else ""
        if ext not in _VIDEO_EXT and ext not in _DOC_EXT:
            skipped += 1
            continue
        try:
            content = zf.read(name)
            path = f"{APP_NAME}/uploads/{user['id']}/{new_id()}.{ext}"
            ct = MIME_TYPES.get(ext, "application/octet-stream")
            result = await asyncio.to_thread(put_object, path, content, ct)
        except Exception as e:
            errors.append(f"{base}: {e}")
            continue
        stored_path = result["path"]
        await db.files.insert_one({
            "id": new_id(), "storage_path": stored_path, "original_filename": base,
            "content_type": ct, "size": result.get("size", len(content)),
            "uploaded_by": user["id"], "is_deleted": False, "created_at": now_iso(),
        })
        url = f"/api/files/{stored_path}"
        order += 1
        title = base.rsplit(".", 1)[0]
        if ext in _VIDEO_EXT:
            doc = {"id": new_id(), "course_id": course_id, "title": title, "description": "Diimpor dari ZIP",
                   "video_type": "upload", "video_url": url, "attachments": [], "order": order, "created_at": now_iso()}
        else:
            doc = {"id": new_id(), "course_id": course_id, "title": title, "description": "Diimpor dari ZIP",
                   "video_type": "document", "video_url": None,
                   "attachments": [{"name": base, "url": url, "content_type": ct}], "order": order, "created_at": now_iso()}
        await db.lessons.insert_one(doc)
        created += 1
    return {"created": created, "skipped": skipped, "errors": errors}
